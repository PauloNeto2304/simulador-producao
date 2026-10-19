"""Gera a planilha modelo (data/template_fabrica.xlsx) que alimenta o simulador.

Todos os números são PLACEHOLDERS realistas — substitua pelos dados da fábrica.

Uso:
    py scripts/gerar_template.py            # não sobrescreve se o arquivo já existir
    py scripts/gerar_template.py --forcar   # sobrescreve
"""
import random
import sys
from datetime import date, time, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

SAIDA = Path(__file__).resolve().parents[1] / "data" / "template_fabrica.xlsx"

CABECALHO = PatternFill("solid", fgColor="1F4E78")
EXEMPLO = PatternFill("solid", fgColor="FFF2CC")  # amarelo = dado de exemplo, substituir

PROCESSOS = [
    (10, "EXT_TUBO", "Extrusão do tubo interno sobre o mandril"),
    (20, "TRANCA", "Aplicação do reforço (trançagem)"),
    (30, "EXT_COBERTURA", "Extrusão da cobertura"),
    (40, "BANDAGEM", "Aplicação da fita de vulcanização e brand mark"),
    (50, "VULCANIZACAO", "Vulcanização na autoclave (lote)"),
    (60, "DESBANDAGEM", "Remoção da fita de vulcanização"),
    (70, "TESTE", "Ejeção do mandril e teste da mangueira"),
]

TIPOS_MAQUINA = [
    "EXTRUSORA", "TRANCADEIRA", "ESPIRALADEIRA", "BANDADEIRA", "AUTOCLAVE", "CABINE_TESTE",
    "BOBINADEIRA_ARAME", "BOBINADEIRA_ESPIRAL", "BOBINADEIRA_TEXTIL", "SOLDA_MANDRIL",
]

MAQUINAS = [
    # maquina, descricao, tipo, decks, jogos_bobina, ativa
    ("EXT-1", "Extrusora 1", "EXTRUSORA", "", "", "S"),
    ("EXT-2", "Extrusora 2", "EXTRUSORA", "", "", "S"),
    ("EXT-3", "Extrusora 3 (mandril)", "EXTRUSORA", "", "", "S"),
    ("TR-01", "Trançadeira 01", "TRANCADEIRA", 1, 2, "S"),
    ("TR-02", "Trançadeira 02", "TRANCADEIRA", 1, 2, "S"),
    ("TR-03", "Trançadeira 03", "TRANCADEIRA", 2, 2, "S"),
    ("TR-04", "Trançadeira 04", "TRANCADEIRA", 2, 2, "S"),
    ("TR-05", "Trançadeira 05", "TRANCADEIRA", 1, 2, "S"),
    ("TR-06", "Trançadeira 06", "TRANCADEIRA", 2, 2, "S"),
    ("TR-07", "Trançadeira 07", "TRANCADEIRA", 2, 2, "S"),
    ("TR-08", "Trançadeira 08", "TRANCADEIRA", 2, 2, "S"),
    ("TR-09", "Trançadeira 09", "TRANCADEIRA", 1, 2, "S"),
    ("TR-10", "Trançadeira 10", "TRANCADEIRA", 1, 2, "S"),
    ("TR-11", "Trançadeira 11", "TRANCADEIRA", 4, 2, "S"),
    ("BAND-1", "Bandadeira 1", "BANDADEIRA", "", "", "S"),
    ("BAND-2", "Bandadeira 2", "BANDADEIRA", "", "", "S"),
    ("AUTOCLAVE", "Autoclave", "AUTOCLAVE", "", "", "S"),
    ("TESTE-01", "Cabine de teste 01", "CABINE_TESTE", "", "", "S"),
    ("TESTE-02", "Cabine de teste 02", "CABINE_TESTE", "", "", "S"),
    ("BOB-A01", "Bobinadeira de arame 01", "BOBINADEIRA_ARAME", "", "", "S"),
    ("BOB-A02", "Bobinadeira de arame 02", "BOBINADEIRA_ARAME", "", "", "S"),
    ("BOB-A03", "Bobinadeira de arame 03", "BOBINADEIRA_ARAME", "", "", "S"),
    ("BOB-A04", "Bobinadeira de arame 04", "BOBINADEIRA_ARAME", "", "", "S"),
    ("BOB-E01", "Bobinadeira de espiral 01", "BOBINADEIRA_ESPIRAL", "", "", "S"),
    ("BOB-E02", "Bobinadeira de espiral 02", "BOBINADEIRA_ESPIRAL", "", "", "S"),
    ("BOB-T01", "Bobinadeira têxtil 01", "BOBINADEIRA_TEXTIL", "", "", "S"),
    ("SOLDA-MAND-01", "Inspeção/solda de mandril 01", "SOLDA_MANDRIL", "", "", "S"),
    ("SOLDA-MAND-02", "Inspeção/solda de mandril 02", "SOLDA_MANDRIL", "", "", "S"),
]

DASH = [(4, 800), (6, 680), (8, 500), (10, 460), (12, 400), (16, 250)]

UM_DECK = ["TR-01", "TR-02", "TR-05", "TR-09", "TR-10"]
DOIS_DECKS = ["TR-03", "TR-04", "TR-06", "TR-07", "TR-08"]

CONSTRUCOES = [
    # construcao, produto, camadas, tipo_fio, tipo_bobinadeira, rendimento_m_por_jogo, tempo_bobinar_jogo_min,
    # fator de tempo (placeholder), trançadeiras possíveis
    ("GH493-6-A", "GH493-6", 1, "ARAME-0.25", "BOBINADEIRA_ARAME", 2500, 150, 0.8, UM_DECK + ["TR-03", "TR-04"]),
    ("GH781-8-A", "GH781-8", 2, "ARAME-0.30", "BOBINADEIRA_ARAME", 1500, 180, 1.0, UM_DECK + DOIS_DECKS),
    ("GH781-8-B", "GH781-8", 2, "ARAME-0.35", "BOBINADEIRA_ARAME", 1300, 180, 0.9, DOIS_DECKS + ["TR-11"]),
    ("GH781-12-A", "GH781-12", 2, "ARAME-0.35", "BOBINADEIRA_ARAME", 1200, 180, 1.2, DOIS_DECKS + ["TR-11"]),
    ("GH781-12-B", "GH781-12", 2, "ARAME-0.40", "BOBINADEIRA_ARAME", 1000, 200, 1.1, ["TR-08", "TR-11", "TR-09", "TR-10"]),
    ("GH900-16-A", "GH900-16", 2, "ARAME-0.50", "BOBINADEIRA_ARAME", 800, 210, 1.6, ["TR-11", "TR-08", "TR-06"]),
]

# seg/metro POR PASSAGEM na trançadeira, conforme o nº de decks da máquina (antes do fator da construção)
TRANCA_BASE = {1: 26.0, 2: 30.0, 4: 24.0}
DECKS = {m: d for m, _, t, d, _, _ in MAQUINAS if t == "TRANCADEIRA"}


def _rota(construcao, fator, trancadeiras, so_ext2=False):
    ext_tubo = [("EXT-2", 0.45)] if so_ext2 else [("EXT-1", 0.40), ("EXT-2", 0.45)]
    etapas = [
        (10, "EXT_TUBO", ext_tubo),
        (20, "TRANCA", [(m, TRANCA_BASE[DECKS[m]]) for m in trancadeiras]),
        (30, "EXT_COBERTURA", [("EXT-1", 0.35), ("EXT-2", 0.38)]),
        (40, "BANDAGEM", [("BAND-1", 0.5), ("BAND-2", 0.5)]),
        (50, "VULCANIZACAO", [("AUTOCLAVE", None)]),
        (60, "DESBANDAGEM", [("BAND-1", 0.4), ("BAND-2", 0.4)]),
        (70, "TESTE", [("TESTE-01", 0.3), ("TESTE-02", 0.3)]),
    ]
    return [
        (construcao, ordem, proc, maq, "" if t is None else round(t * fator, 2))
        for ordem, proc, alternativas in etapas
        for maq, t in alternativas
    ]


ROTAS = [linha for c, _, _, _, _, _, _, f, tr in CONSTRUCOES for linha in _rota(c, f, tr, so_ext2=c.startswith("GH900"))]

SETUPS = (
    [("EXT-1", 15, 45), ("EXT-2", 15, 45)]
    + [(m, 10, 30) for m in DECKS]
    + [("BAND-1", 5, 10), ("BAND-2", 5, 10), ("TESTE-01", 5, 15), ("TESTE-02", 5, 15)]
)

PARAMETROS = [
    ("alimentacao_fio_min_por_deck", 47, "Tempo de alimentação de fio por deck (2 decks = 2x)"),
    ("troca_tubo_min", 15, "Troca de tubo na trançadeira (a cada carretel e a cada passagem)"),
    ("pulmao_tubos_por_trancadeira", 2, "Máx. de tubos prontos esperando cada trançadeira (a 'corda')"),
    ("inspecao_solda_mandril_min", 20, "Tempo de inspeção/solda de um mandril após a ejeção"),
]

MANDRIS = [(4, 40), (6, 40), (8, 50), (10, 30), (12, 40), (16, 25)]

AUTOCLAVE = [
    ("capacidade_carreteis", 5, "Máximo de carretéis por ciclo"),
    ("ciclo_min", 120, "Ciclo padrão (usado quando o produto não está na aba autoclave_produtos)"),
    ("carga_descarga_min", 15, "Tempo para carregar + descarregar"),
    ("espera_max_lote_min", 60, "Tempo máximo esperando completar o lote antes de rodar incompleto"),
]

AUTOCLAVE_PRODUTOS = [
    ("GH781-8", 120, "GERAL"),
    ("GH781-12", 120, "GERAL"),
    ("GH493-6", 120, "GERAL"),
    ("GH900-16", 150, "GH900"),
]

TURNOS = [
    ("T1", time(5, 20), time(13, 57), "seg-sab"),
    ("T2", time(13, 57), time(22, 0), "seg-sab"),
    ("T3", time(22, 0), time(5, 20), "seg-sab"),
]


def _pedidos(n=40, semente=42):
    rnd = random.Random(semente)
    produtos = ["GH781-8", "GH781-8", "GH781-12", "GH493-6", "GH493-6", "GH900-16"]
    linhas = []
    for i in range(1, n + 1):
        produto = rnd.choice(produtos)
        metros = rnd.choice([1000, 1500, 2000, 3000, 4000, 6000])
        entrega = date(2026, 11, 4) + timedelta(days=rnd.randint(0, 14))
        linhas.append((f"PED-{i:03d}", produto, metros, entrega, ""))
    return linhas


PEDIDOS = _pedidos()

LEIA_ME = [
    "PLANILHA DE ENTRADA DO SIMULADOR DE PRODUÇÃO",
    "",
    "Células em AMARELO são exemplos (placeholders): substitua pelos dados reais (pode apagar/adicionar linhas).",
    "Não renomeie abas nem cabeçalhos.",
    "",
    "processos   - etapas da linha principal, na ordem (coluna 'ordem').",
    "maquinas    - uma linha por máquina física (inclusive bobinadeiras e solda de mandril). ativa = S/N.",
    "              decks e jogos_bobina: só para trançadeiras (jogos = carrinhos próprios da máquina, não emprestam).",
    "dash        - metros padrão por carretel de cada dash/bitola. Pedidos são arredondados para carretéis cheios.",
    "construcoes - cada produto pode ter várias construções; o simulador escolhe qual usar.",
    "              camadas: nº de camadas de trança. Passagens na trançadeira = camadas / decks da máquina",
    "              (ex.: 2 camadas em máquina de 1 deck = 2 passagens).",
    "              tipo_fio: trocar de tipo de fio descarta o resto do jogo montado (refugo).",
    "              rendimento_m_por_jogo: metros de mangueira (por passagem) que um jogo de bobinas rende.",
    "              tempo_bobinar_jogo_min: tempo de UMA bobinadeira para bobinar um jogo inteiro.",
    "rotas       - uma linha por (construção, etapa, máquina possível) com o tempo em segundos/metro.",
    "              TRANCA: tempo POR PASSAGEM. Várias linhas na mesma etapa = máquinas alternativas.",
    "              VULCANIZACAO: deixe seg_por_metro vazio (o tempo vem das abas de autoclave).",
    "setups      - minutos de setup por máquina ao trocar de produto (mesmo dash) e ao trocar de dash.",
    "parametros  - regras da trançadeira e do mandril.",
    "mandris     - quantidade de mandris disponíveis por dash (ciclo fechado: voltam após a ejeção).",
    "autoclave   - parâmetros do lote.",
    "autoclave_produtos - ciclo (min) e grupo_mistura: só carretéis do MESMO grupo vulcanizam juntos.",
    "pedidos     - carteira: produto, metros e data de entrega. O dash é lido do sufixo do nome (GH781-8 -> 8).",
    "turnos      - horários de trabalho. dias: 'seg-sab', 'seg-sex', 'seg,qua,sex'...",
]


def escrever_aba(wb, nome, cabecalho, linhas, larguras):
    ws = wb.create_sheet(nome)
    ws.append(cabecalho)
    for cel in ws[1]:
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = CABECALHO
        cel.alignment = Alignment(horizontal="center")
    for linha in linhas:
        ws.append(list(linha))
        for cel in ws[ws.max_row]:
            cel.fill = EXEMPLO
    for i, largura in enumerate(larguras, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = largura
    ws.freeze_panes = "A2"
    return ws


def main():
    if SAIDA.exists() and "--forcar" not in sys.argv:
        print(f"{SAIDA} já existe (use --forcar para sobrescrever).")
        return

    wb = Workbook()
    leia = wb.active
    leia.title = "LEIA-ME"
    for texto in LEIA_ME:
        leia.append([texto])
    leia["A1"].font = Font(bold=True, size=14)
    leia.column_dimensions["A"].width = 120

    escrever_aba(wb, "processos", ["ordem", "processo", "descricao"], PROCESSOS, [8, 18, 50])
    maquinas = escrever_aba(
        wb, "maquinas", ["maquina", "descricao", "tipo", "decks", "jogos_bobina", "ativa"], MAQUINAS,
        [16, 32, 22, 8, 13, 8],
    )
    escrever_aba(wb, "dash", ["dash", "metros_por_carretel"], DASH, [8, 22])
    escrever_aba(
        wb, "construcoes",
        ["construcao", "produto", "camadas", "tipo_fio", "tipo_bobinadeira", "rendimento_m_por_jogo",
         "tempo_bobinar_jogo_min"],
        [c[:7] for c in CONSTRUCOES], [16, 14, 10, 14, 22, 22, 24],
    )
    rotas = escrever_aba(
        wb, "rotas", ["construcao", "ordem", "processo", "maquina", "seg_por_metro"], ROTAS, [16, 8, 18, 16, 15]
    )
    escrever_aba(
        wb, "setups", ["maquina", "setup_troca_produto_min", "setup_troca_dash_min"], SETUPS, [16, 26, 24]
    )
    escrever_aba(wb, "parametros", ["parametro", "valor", "descricao"], PARAMETROS, [30, 10, 70])
    escrever_aba(wb, "mandris", ["dash", "quantidade"], MANDRIS, [8, 12])
    escrever_aba(wb, "autoclave", ["parametro", "valor", "descricao"], AUTOCLAVE, [24, 10, 70])
    escrever_aba(
        wb, "autoclave_produtos", ["produto", "ciclo_min", "grupo_mistura"], AUTOCLAVE_PRODUTOS, [16, 10, 16]
    )
    pedidos = escrever_aba(
        wb, "pedidos", ["pedido", "produto", "metros", "data_entrega", "prioridade"], PEDIDOS, [12, 16, 10, 14, 11]
    )
    turnos = escrever_aba(wb, "turnos", ["turno", "inicio", "fim", "dias"], TURNOS, [8, 8, 8, 12])

    listas = wb.create_sheet("listas")
    listas.append(["tipo_maquina"])
    for t in TIPOS_MAQUINA:
        listas.append([t])
    listas.sheet_state = "hidden"

    dv_tipo = DataValidation(type="list", formula1="=listas!$A$2:$A$50", allow_blank=False)
    dv_sn = DataValidation(type="list", formula1='"S,N"', allow_blank=False)
    maquinas.add_data_validation(dv_tipo)
    maquinas.add_data_validation(dv_sn)
    dv_tipo.add("C2:C500")
    dv_sn.add("F2:F500")

    dv_proc = DataValidation(type="list", formula1="=processos!$B$2:$B$50", allow_blank=False)
    dv_maq = DataValidation(type="list", formula1="=maquinas!$A$2:$A$500", allow_blank=False)
    rotas.add_data_validation(dv_proc)
    rotas.add_data_validation(dv_maq)
    dv_proc.add("C2:C5000")
    dv_maq.add("D2:D5000")

    for linha in range(2, 5001):
        pedidos.cell(linha, 4).number_format = "dd/mm/yyyy"
    for linha in range(2, 50):
        for col in (2, 3):
            turnos.cell(linha, col).number_format = "hh:mm"

    SAIDA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SAIDA)
    print(f"Planilha gerada em {SAIDA}")


if __name__ == "__main__":
    main()
